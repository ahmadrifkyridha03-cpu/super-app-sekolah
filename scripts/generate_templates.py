import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

os.makedirs("d:/test/public/templates", exist_ok=True)
os.makedirs("d:/test/templates", exist_ok=True)

# Styling Definitions (Institutional Navy & Academic Gold)
HEADER_FILL = PatternFill(start_color="0F2942", end_color="0F2942", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="0F2942")
SUBTITLE_FONT = Font(name="Calibri", size=10, italic=True, color="555555")
DATA_FONT = Font(name="Calibri", size=10)
NOTE_FONT = Font(name="Calibri", size=9, color="777777")

THIN_BORDER = Border(
    left=Side(style='thin', color='D0D7DE'),
    right=Side(style='thin', color='D0D7DE'),
    top=Side(style='thin', color='D0D7DE'),
    bottom=Side(style='thin', color='D0D7DE')
)

def auto_fit_columns(ws, min_col=1, max_col=None):
    if max_col is None:
        max_col = ws.max_column
    for col in range(min_col, max_col + 1):
        col_letter = get_column_letter(col)
        max_len = 0
        for cell in ws[col_letter]:
            val = str(cell.value or '')
            if cell.row < 3:
                continue
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

# ==============================================================================
# 1. TEMPLATE DATA SISWA
# ==============================================================================
wb_siswa = openpyxl.Workbook()

# Sheet Data
ws_s = wb_siswa.active
ws_s.title = "DATA_SISWA"

# Title header
ws_s.merge_cells("A1:K1")
ws_s["A1"] = "TEMPLATE IMPORT DATA MASTER SISWA - SMA NEGERI UNGGULAN BANGSA"
ws_s["A1"].font = TITLE_FONT
ws_s["A1"].alignment = Alignment(vertical="center")

ws_s.merge_cells("A2:K2")
ws_s["A2"] = "*Petunjuk: Jangan mengubah nama atau urutan kolom baris 4. Format tanggal lahir: YYYY-MM-DD. Kolom bertanda (*) wajib diisi."
ws_s["A2"].font = SUBTITLE_FONT

headers_siswa = [
    "No", "NIS*", "NISN*", "Nama Lengkap*", "Jenis Kelamin (L/P)*", 
    "Tempat Lahir*", "Tanggal Lahir (YYYY-MM-DD)*", "Kode Kelas*", 
    "Nama Ayah", "Nama Ibu", "No. Telepon / WhatsApp Ortu*"
]

for col_num, header in enumerate(headers_siswa, 1):
    c = ws_s.cell(row=4, column=col_num)
    c.value = header
    c.fill = HEADER_FILL
    c.font = HEADER_FONT
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = THIN_BORDER
ws_s.row_dimensions[4].height = 28

sample_siswa = [
    [1, "24251001", "0081234567", "Ahmad Fauzi Pratama", "L", "Jakarta", "2008-04-12", "X-A", "Bambang Pratama", "Siti Rahmawati", "081234567890"],
    [2, "24251002", "0082345678", "Annisa Putri Maharani", "P", "Bandung", "2008-07-25", "X-A", "Dedi Kurniawan", "Endang Susilowati", "081234567891"],
    [3, "24251003", "0083456789", "Bagas Danendra", "L", "Surabaya", "2008-01-30", "X-A", "Gunawan Prasetyo", "Dewi Lestari", "081234567892"],
    [4, "24251004", "0084567890", "Cantika Aulia Azzahra", "P", "Semarang", "2008-09-18", "X-B", "Hendra Wijaya", "Ratna Sari", "081234567893"],
    [5, "24251005", "0085678901", "Dimas Arya Nugraha", "L", "Yogyakarta", "2008-11-05", "X-B", "Irwan Nugraha", "Sri Wahyuni", "081234567894"],
]

for r_idx, row_data in enumerate(sample_siswa, 5):
    for c_idx, val in enumerate(row_data, 1):
        cell = ws_s.cell(row=r_idx, column=c_idx, value=val)
        cell.font = DATA_FONT
        cell.border = THIN_BORDER
        if c_idx in [1, 2, 3, 5, 7, 8]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(vertical="center")
    ws_s.row_dimensions[r_idx].height = 20

auto_fit_columns(ws_s)

# Sheet Petunjuk
ws_petunjuk_s = wb_siswa.create_sheet(title="PETUNJUK_PENGISIAN")
ws_petunjuk_s["A1"] = "PANDUAN LENGKAP PENGISIAN TEMPLATE SISWA"
ws_petunjuk_s["A1"].font = TITLE_FONT

petunjuk_siswa_data = [
    ("Kolom", "Format Data", "Contoh Valid", "Keterangan"),
    ("NIS", "Angka / Karakter Teks (Unik)", "24251001", "Nomor Induk Siswa lokal sekolah. Tidak boleh ada yang sama."),
    ("NISN", "10 Digit Angka Tepat (Unik)", "0081234567", "Nomor Induk Siswa Nasional resmi Kemdikbudristek."),
    ("Nama Lengkap", "Teks Huruf Kapital/Sesuai Ijazah", "Ahmad Fauzi Pratama", "Nama resmi siswa sesuai Akta Kelahiran/Ijazah SMP."),
    ("Jenis Kelamin", "Tepat 'L' atau 'P'", "L", "L = Laki-laki, P = Perempuan."),
    ("Tempat Lahir", "Teks Nama Kota/Kabupaten", "Jakarta", "Kota kelahiran sesuai akta."),
    ("Tanggal Lahir", "Format YYYY-MM-DD", "2008-04-12", "Tahun 4 digit, Bulan 2 digit, Tanggal 2 digit."),
    ("Kode Kelas", "Teks Kode Rombel Resmi", "X-A", "Harus cocok dengan kode kelas master di sistem."),
    ("Nama Ayah / Ibu", "Teks Nama Orang Tua", "Bambang Pratama", "Nama orang tua kandung atau wali."),
    ("No. Telepon / WA", "Format 08xxxxxxxxxx / 628xxxxxxxxxx", "081234567890", "Digunakan untuk notifikasi absensi dan e-raport.")
]

for r, row in enumerate(petunjuk_siswa_data, 3):
    for c, val in enumerate(row, 1):
        cell = ws_petunjuk_s.cell(row=r, column=c, value=val)
        cell.border = THIN_BORDER
        if r == 3:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.font = DATA_FONT
            cell.alignment = Alignment(vertical="center")
    ws_petunjuk_s.row_dimensions[r].height = 22

auto_fit_columns(ws_petunjuk_s)

wb_siswa.save("d:/test/public/templates/template_siswa.xlsx")
wb_siswa.save("d:/test/templates/template_siswa.xlsx")
print("Template siswa created successfully.")

# ==============================================================================
# 2. TEMPLATE DATA GURU
# ==============================================================================
wb_guru = openpyxl.Workbook()

ws_g = wb_guru.active
ws_g.title = "DATA_GURU"

# Title header
ws_g.merge_cells("A1:J1")
ws_g["A1"] = "TEMPLATE IMPORT DATA MASTER GURU & TENAGA PENDIDIK - SMA NEGERI UNGGULAN BANGSA"
ws_g["A1"].font = TITLE_FONT
ws_g["A1"].alignment = Alignment(vertical="center")

ws_g.merge_cells("A2:J2")
ws_g["A2"] = "*Petunjuk: Jangan mengubah nama atau urutan kolom baris 4. Kolom NIP boleh dikosongkan jika Guru Honorer/Non-PNS. Kolom (*) wajib diisi."
ws_g["A2"].font = SUBTITLE_FONT

headers_guru = [
    "No", "NIP / NUPTK", "Nama Lengkap & Gelar*", "Email Resmi / Akun*", 
    "Jenis Kelamin (L/P)*", "Mata Pelajaran Utama*", "Mata Pelajaran Tambahan", 
    "Jabatan / Penugasan*", "Wali Kelas di Rombel", "No. Telepon / WhatsApp*"
]

for col_num, header in enumerate(headers_guru, 1):
    c = ws_g.cell(row=4, column=col_num)
    c.value = header
    c.fill = HEADER_FILL
    c.font = HEADER_FONT
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = THIN_BORDER
ws_g.row_dimensions[4].height = 28

sample_guru = [
    [1, "198002152005011004", "Drs. Budi Santoso, M.Pd.", "budi.santoso@sekolah.sch.id", "L", "Matematika", "", "Wali Kelas", "X-A", "081298765001"],
    [2, "198406122008012009", "Siti Aminah, S.Pd., M.Si.", "siti.aminah@sekolah.sch.id", "P", "Bahasa Indonesia", "", "Wali Kelas", "X-B", "081298765002"],
    [3, "197811032003121002", "Dr. Irwan Setiawan, M.Kom.", "irwan.setiawan@sekolah.sch.id", "L", "Informatika", "Fisika", "Wakil Kepala Sekolah Bidang Kurikulum", "", "081298765003"],
    [4, "199008242015022003", "Dewi Anggraini, S.Pd.", "dewi.anggraini@sekolah.sch.id", "P", "Bahasa Inggris", "", "Guru Mata Pelajaran", "", "081298765004"],
    [5, "-", "Fajar Pratama, S.Pd.", "fajar.pratama@sekolah.sch.id", "L", "Pendidikan Jasmani & Olahraga", "", "Guru Mata Pelajaran", "", "081298765005"],
]

for r_idx, row_data in enumerate(sample_guru, 5):
    for c_idx, val in enumerate(row_data, 1):
        cell = ws_g.cell(row=r_idx, column=c_idx, value=val)
        cell.font = DATA_FONT
        cell.border = THIN_BORDER
        if c_idx in [1, 2, 5, 9]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(vertical="center")
    ws_g.row_dimensions[r_idx].height = 20

auto_fit_columns(ws_g)

# Sheet Petunjuk Guru
ws_petunjuk_g = wb_guru.create_sheet(title="PETUNJUK_PENGISIAN")
ws_petunjuk_g["A1"] = "PANDUAN LENGKAP PENGISIAN TEMPLATE GURU"
ws_petunjuk_g["A1"].font = TITLE_FONT

petunjuk_guru_data = [
    ("Kolom", "Format Data", "Contoh Valid", "Keterangan"),
    ("NIP / NUPTK", "18 Digit Angka (PNS) atau '-' jika Honorer", "198002152005011004", "Nomor Induk Pegawai resmi. Jika belum ada isi '-'."),
    ("Nama Lengkap & Gelar", "Teks Nama beserta Gelar Akademik", "Drs. Budi Santoso, M.Pd.", "Nama yang akan dicetak di e-raport dan jadwal mengajar."),
    ("Email Resmi / Akun", "Format Email Valid (Unik)", "budi.santoso@sekolah.sch.id", "Digunakan sebagai username login utama guru."),
    ("Jenis Kelamin", "Tepat 'L' atau 'P'", "L", "L = Laki-laki, P = Perempuan."),
    ("Mata Pelajaran Utama", "Teks Nama Mata Pelajaran", "Matematika", "Mata pelajaran induk yang diampu."),
    ("Jabatan / Penugasan", "Teks Penugasan", "Wali Kelas", "Contoh: Guru Mata Pelajaran, Wali Kelas, Waka Kurikulum, Kepala Laboratorium."),
    ("Wali Kelas di Rombel", "Kode Kelas (Jika menjabat)", "X-A", "Kosongkan jika bukan wali kelas."),
    ("No. Telepon / WA", "Format 08xxxxxxxxxx", "081298765001", "Nomor kontak untuk koordinasi sekolah.")
]

for r, row in enumerate(petunjuk_guru_data, 3):
    for c, val in enumerate(row, 1):
        cell = ws_petunjuk_g.cell(row=r, column=c, value=val)
        cell.border = THIN_BORDER
        if r == 3:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.font = DATA_FONT
            cell.alignment = Alignment(vertical="center")
    ws_petunjuk_g.row_dimensions[r].height = 22

auto_fit_columns(ws_petunjuk_g)

wb_guru.save("d:/test/public/templates/template_guru.xlsx")
wb_guru.save("d:/test/templates/template_guru.xlsx")
print("Template guru created successfully.")
